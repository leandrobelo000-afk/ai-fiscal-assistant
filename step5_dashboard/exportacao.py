"""
step5_dashboard/exportacao.py
-----------------------------
Gera o Excel (abas Resumo e Itens, mesmo layout da Etapa 3) a partir da base SQLite.
"""

import io

import pandas as pd
from openpyxl.utils import get_column_letter

from step3_spreadsheet.writer import COR_CABECALHO, COR_CABECALHO_2, _estilizar_cabecalho


def _formatar_cnpj(cnpj) -> str | None:
    if not isinstance(cnpj, str) or len(cnpj) != 14:
        return cnpj
    return f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"


def exportar_excel(notas: pd.DataFrame, itens: pd.DataFrame) -> bytes:
    """Devolve o arquivo .xlsx em memória, pronto para st.download_button."""
    qtd_itens = itens.groupby("nota_id").size()
    resumo = pd.DataFrame({
        "Data Processamento": pd.to_datetime(notas["data_processamento"]).dt.strftime("%d/%m/%Y %H:%M"),
        "Número NF": notas["numero"],
        "Data Emissão": notas["data_emissao"].dt.strftime("%d/%m/%Y"),
        "CNPJ Emitente": notas["cnpj_emitente"].map(_formatar_cnpj),
        "Nome Emitente": notas["nome_emitente"],
        "Valor Total (R$)": notas["valor_total"],
        "Impostos (R$)": notas["impostos"],
        "Qtd Itens": notas["id"].map(qtd_itens).fillna(0).astype(int),
    })
    aba_itens = pd.DataFrame({
        "Número NF": itens["numero"],
        "Data Emissão": itens["data_emissao"].dt.strftime("%d/%m/%Y"),
        "Nome Emitente": itens["nome_emitente"],
        "Descrição Item": itens["descricao"],
        "Valor Item (R$)": itens["valor"],
    })

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as excel:
        for nome, df, cor in (("Resumo", resumo, COR_CABECALHO), ("Itens", aba_itens, COR_CABECALHO_2)):
            df.to_excel(excel, sheet_name=nome, index=False)
            ws = excel.sheets[nome]
            _estilizar_cabecalho(ws, list(df.columns), cor)
            for col_idx, coluna in enumerate(df.columns, start=1):
                if "(R$)" in coluna:
                    for celula in ws[get_column_letter(col_idx)][1:]:
                        celula.number_format = '"R$" #,##0.00'
    return buffer.getvalue()
