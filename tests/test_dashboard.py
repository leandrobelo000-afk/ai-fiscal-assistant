"""Renderiza cada página do dashboard com uma base de exemplo e verifica que não há erro."""
import os
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
import io

from step5_dashboard import database
from step5_dashboard.exportacao import exportar_excel

PASTA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "step5_dashboard", "paginas")


def _popular(banco):
    for i, (data, cnpj, nome, valor) in enumerate([
        ("10/07/2026", "11.111.111/0001-11", "Posto A", 700.0),
        ("20/08/2026", "22.222.222/0001-22", "Mercado B", 200.0),
        ("05/09/2026", "11.111.111/0001-11", "Posto A", 80.0),
    ], start=1):
        database.salvar_nota(SimpleNamespace(
            numero=str(i), data_emissao=data, cnpj_emitente=cnpj, nome_emitente=nome,
            valor_total=valor, impostos=valor * 0.1, itens=[{"descricao": "Item", "valor": valor}],
            rota="vision", confianca="alta"), banco)


@pytest.fixture
def na_pasta_com_banco(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _popular(database.CAMINHO_BANCO)


@pytest.mark.parametrize("pagina", ["visao_geral.py", "fornecedores.py", "itens.py"])
def test_paginas_renderizam(na_pasta_com_banco, pagina):
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(os.path.join(PASTA, pagina), default_timeout=30).run()
    assert not at.exception
    assert at.title[0].value


def test_visao_geral_mostra_kpis(na_pasta_com_banco):
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(os.path.join(PASTA, "visao_geral.py"), default_timeout=30).run()
    assert at.metric[0].value == "R$ 980,00"
    assert at.metric[1].value == "3"


@pytest.mark.parametrize("pagina", ["visao_geral.py", "fornecedores.py", "itens.py"])
def test_paginas_com_base_vazia(tmp_path, monkeypatch, pagina):
    monkeypatch.chdir(tmp_path)
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(os.path.join(PASTA, pagina), default_timeout=30).run()
    assert not at.exception
    assert at.info


def test_exportar_excel(na_pasta_com_banco):
    conteudo = exportar_excel(database.carregar_notas(), database.carregar_itens())
    wb = load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == ["Resumo", "Itens"]
    resumo = wb["Resumo"]
    assert resumo["D2"].value == "11.111.111/0001-11"
    assert resumo.max_row == 4


def test_pagina_processar_nota_sem_tesseract(tmp_path, monkeypatch):
    """A tela da Etapa 4 abre mesmo quando o Tesseract não pode ser instalado."""
    monkeypatch.chdir(tmp_path)
    from step1_ocr import extractor
    monkeypatch.setattr(extractor, "_localizar_tesseract", lambda: None)
    monkeypatch.setattr(extractor.platform, "system", lambda: "Linux")
    from streamlit.testing.v1 import AppTest
    app = os.path.join(os.path.dirname(PASTA), "..", "step4_interface", "app.py")
    at = AppTest.from_file(app, default_timeout=30).run()
    assert not at.exception
    assert "Tesseract" in at.warning[0].value
