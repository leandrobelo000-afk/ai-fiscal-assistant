from types import SimpleNamespace

from openpyxl import load_workbook

from step3_spreadsheet import writer


def _nota(numero, cnpj, valor):
    return SimpleNamespace(numero=numero, data_emissao="05/10/2026", cnpj_emitente=cnpj,
                           nome_emitente="Fornecedor", valor_total=valor, impostos=0.0,
                           itens=[{"descricao": "x", "valor": valor}])


def test_total_geral_fora_dos_dados(tmp_path):
    planilha = str(tmp_path / "n.xlsx")
    writer.gravar_nota(_nota("1", "11.111.111/0001-11", 10.0), planilha)
    writer.gravar_nota(_nota("2", "11.111.111/0001-11", 20.0), planilha)
    ws = load_workbook(planilha)["Resumo"]
    rotulos = [ws.cell(r, 5).value for r in range(2, ws.max_row + 1)]
    assert "TOTAL GERAL" not in rotulos
    assert ws.max_row == 3
    assert ws.cell(2, writer.COLUNA_TOTAL).value == "=SUM(F:F)"


def test_remove_totais_de_planilhas_antigas(tmp_path):
    planilha = str(tmp_path / "n.xlsx")
    writer.criar_planilha(planilha)
    wb = load_workbook(planilha)
    ws = wb["Resumo"]
    ws.append(["", "1", "", "", "A", 10.0, 0, 1])
    ws.append([None] * 8)
    ws.append([None, None, None, None, "TOTAL GERAL", "=SUM(F2:F2)"])
    wb.save(planilha)
    writer.gravar_nota(_nota("2", "22.222.222/0001-22", 20.0), planilha)
    ws = load_workbook(planilha)["Resumo"]
    assert [ws.cell(r, 6).value for r in range(2, ws.max_row + 1)] == [10.0, 20.0]


def test_duplicata_considera_cnpj(tmp_path):
    planilha = str(tmp_path / "n.xlsx")
    assert writer.gravar_nota(_nota("1", "11.111.111/0001-11", 10.0), planilha)
    assert not writer.gravar_nota(_nota("1", "11111111000111", 10.0), planilha)
    assert writer.gravar_nota(_nota("1", "22.222.222/0001-22", 10.0), planilha)
